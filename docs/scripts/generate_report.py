"""Generate an editable Word report from docs/BAO_CAO_DO_AN.md.
Does not read secrets or modify application source files.
"""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "BAO_CAO_DO_AN.md"
OUTPUT = ROOT / "docs" / "BAO_CAO_DO_AN.docx"

def set_cell_text(cell, text, bold=False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text.strip())
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

def add_field(paragraph, instruction):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)

def set_font(style, size=13, bold=False):
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold

def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
    normal = doc.styles["Normal"]
    set_font(normal, 13)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for level, size in [(1, 16), (2, 14), (3, 13)]:
        style = doc.styles[f"Heading {level}"]
        set_font(style, size, True)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(8)
    for name in ("Table Caption", "Figure Caption"):
        if name not in [s.name for s in doc.styles]:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_font(doc.styles[name], 11)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("Báo cáo đồ án — Mộc Café").italic = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Trang ")
    add_field(footer, "PAGE")

    # Cover page
    for _ in range(5): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO ĐỒ ÁN WEB APPLICATION"); r.bold = True; r.font.size = Pt(18)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HỆ THỐNG ĐẶT MÓN MỘC CAFÉ"); r.bold = True; r.font.size = Pt(20)
    doc.add_paragraph("\nTrường: [CHƯA CÓ DỮ LIỆU]\nKhoa/Bộ môn: [CHƯA CÓ DỮ LIỆU]\nSinh viên: [CHƯA CÓ DỮ LIỆU]\nMSSV: [CHƯA CÓ DỮ LIỆU]\nGiảng viên: [CHƯA CÓ DỮ LIỆU]")
    p = doc.add_paragraph("[Địa điểm], 2026"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    p = doc.add_paragraph("MỤC LỤC", style="Heading 1")
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()

    correction = ROOT / "docs" / "DIEU_CHINH_BAO_CAO.md"
    source_text = SOURCE.read_text(encoding="utf-8")
    if correction.exists():
        source_text += "\n\n# Phụ lục D. Đối chiếu và điều chỉnh source\n\n" + correction.read_text(encoding="utf-8")
    lines = source_text.splitlines()
    in_code, table_rows = False, []
    def flush_table():
        nonlocal table_rows
        if len(table_rows) >= 2:
            rows = [r for i, r in enumerate(table_rows) if i != 1]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, value in enumerate(row): set_cell_text(table.cell(i, j), value, i == 0)
        table_rows = []
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|") and line.endswith("|"):
            table_rows.append([x.strip() for x in line.strip("|").split("|")]); continue
        flush_table()
        if line.startswith("```"):
            in_code = not in_code; continue
        if in_code:
            p = doc.add_paragraph(style="Normal"); p.paragraph_format.left_indent = Cm(0.8)
            run = p.add_run(line); run.font.name = "Consolas"; run.font.size = Pt(9); continue
        if not line.strip(): continue
        if line.startswith("### "): doc.add_heading(line[4:], level=3); continue
        if line.startswith("## "): doc.add_heading(line[3:], level=2); continue
        if line.startswith("# "): doc.add_heading(line[2:], level=1); continue
        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet"); continue
        if line[:3].isdigit() and ". " in line[:5]: doc.add_paragraph(line, style="List Number"); continue
        p = doc.add_paragraph(line.replace("**", ""))
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    flush_table()
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == "__main__": main()
